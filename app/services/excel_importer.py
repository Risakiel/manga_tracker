"""Import the 'Suivis Manga' Excel export into the Manga table.

Non-destructive: existing rows (matched by MangaUpdates URL, falling back to
server folder name) only have their Excel-owned fields refreshed
(category/title/folder/baseline chapter count/raw status). Fields owned by the
sync services (mu_latest_chapter, author, genres, cover, ...) are left alone so
a re-import never clobbers data fetched from MangaUpdates/AniList/Suwayomi.
"""

import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import openpyxl
from sqlmodel import Session, select

from app.models import Category, Manga, Status, SyncLog, SyncSource, SyncStatus
from app.services.mangaupdates_client import extract_series_id

SHEET_NAME = "LISTS"
_LEADING_NUMBER_RE = re.compile(r"-?\d+(?:\.\d+)?")

_STATUS_MAP = {
    "terminé (complete)": Status.complete,
    "en cours (ongoing)": Status.ongoing,
    "en pause (hiatus)": Status.hiatus,
    "en pause / abandonné (hiatus/dropped)": Status.hiatus,
    "abandonné (cancelled)": Status.cancelled,
}


def _map_status(raw: str) -> Status:
    return _STATUS_MAP.get((raw or "").strip().lower(), Status.unknown)


def _parse_last_scan(value) -> Optional[float]:
    """Most rows hold a plain chapter count, but some hold e.g. '20 V' (20 volumes)."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = _LEADING_NUMBER_RE.search(str(value))
    return float(match.group()) if match else None


@dataclass
class ImportSummary:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)


def import_excel(path: Path, session: Session) -> ImportSummary:
    workbook = openpyxl.load_workbook(path, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"expected a '{SHEET_NAME}' sheet, found: {workbook.sheetnames}")
    sheet = workbook[SHEET_NAME]

    summary = ImportSummary()

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell is None for cell in row):
            continue
        libraries, title_en, server_folder, mu_url, last_scan, status_label = (list(row) + [None] * 6)[:6]

        if not title_en or not mu_url:
            summary.skipped += 1
            summary.errors.append(f"row {row_idx}: missing title or MangaUpdates URL, skipped")
            continue

        existing = session.exec(select(Manga).where(Manga.mangaupdates_url == mu_url)).first()
        if existing is None and server_folder:
            existing = session.exec(select(Manga).where(Manga.server_folder == server_folder)).first()

        category = Category.pornhwa if (libraries or "").strip().lower() == "pornhwa" else Category.manga

        if existing is None:
            manga = Manga(
                category=category,
                title_en=title_en,
                server_folder=server_folder or title_en,
                mangaupdates_url=mu_url,
                mangaupdates_id=extract_series_id(mu_url),
                status_raw=status_label or "",
                status=_map_status(status_label or ""),
                baseline_chapter_count=_parse_last_scan(last_scan),
            )
            session.add(manga)
            summary.created += 1
        else:
            existing.category = category
            existing.title_en = title_en
            existing.server_folder = server_folder or existing.server_folder
            existing.mangaupdates_url = mu_url
            if existing.mangaupdates_id is None:
                existing.mangaupdates_id = extract_series_id(mu_url)
            existing.status_raw = status_label or existing.status_raw
            parsed_last_scan = _parse_last_scan(last_scan)
            existing.baseline_chapter_count = (
                parsed_last_scan if parsed_last_scan is not None else existing.baseline_chapter_count
            )
            existing.updated_at = datetime.now(timezone.utc)
            session.add(existing)
            summary.updated += 1

    session.add(
        SyncLog(
            manga_id=None,
            source=SyncSource.excel,
            status=SyncStatus.error if summary.errors else SyncStatus.success,
            message=f"created={summary.created} updated={summary.updated} skipped={summary.skipped}",
        )
    )
    session.commit()
    return summary
